"""Salida para revisión manual.

    CON ERROR/<EPS>/
        REPORTE <EPS>.xlsx        resumen, discrepancias (valor por fuente) y facturas
        <FACTURA>/                copia literal de la carpeta original
            _REVISION.txt         qué campo difiere, valor en cada documento y cuál es el sospechoso

Solo se copian facturas con al menos una discrepancia (campo, línea o hallazgo aviso/grave).
Los originales no se tocan. Contiene datos de pacientes: la carpeta no va al repositorio.
"""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .explicar import explicar
from .base.schemas import ResultadoFactura

MARCA = ".generado_por_validador"

NOMBRE_CAMPO = {
    "numero_factura": "Número de factura",
    "cufe": "CUFE",
    "nit_emisor": "NIT emisor",
    "nit_eps": "NIT EPS",
    "paciente_id": "Documento del paciente",
    "paciente_nombre": "Nombre del paciente",
    "numero_autorizacion": "Número de autorización",
    "id_interno_paciente": "Identificador interno",
    "valor_total": "Valor total",
    "fecha_ingreso": "Fecha de ingreso",
    "fecha_egreso": "Fecha de egreso",
    "fecha_emision": "Fecha de emisión",
}
NOMBRE_HALLAZGO = {
    "hc": "Historia clínica",
    "documentos": "Documentos",
    "dian": "DIAN",
    "cuv": "Validación Ministerio (CUV)",
    "duplicado": "Duplicado",
    "extraccion": "Lectura de documento",
    "ocr": "Autorización escaneada no legible",
    "procesamiento": "Error al procesar la factura",
    "regla": "Error en una regla de validación",
}
# (clave de fuente, encabezado de columna, dónde mirar)
FUENTES = [
    ("rips", "RIPS (JSON)", "FE….json"),
    ("xml", "XML factura", "FE….xml"),
    ("xml_nota", "Nota del XML", "FE….xml, cbc:Note"),
    ("pdf_factura", "PDF factura", "FEV/FVS"),
    ("hc", "Historia clínica", "HEV"),
    ("autorizacion_digital", "Encabezado de la autorización (hospital)", "PDE, encabezado"),
    ("autorizacion_scan", "Autorización de la EPS (escaneada)", "PDE, imagen"),
    ("cuv", "CUV", "FE…_CUV.json"),
    ("carpeta", "Carpeta", "nombre de la carpeta"),
]
NOMBRE_FUENTE = {k: v for k, v, _ in FUENTES}


# --- discrepancias -------------------------------------------------------------


def discrepancias(r: ResultadoFactura) -> list[dict]:
    """Una fila por cada campo, línea o hallazgo con resultado aviso o grave."""
    filas: list[dict] = []
    for c in r.campos:
        if c.resultado in ("aviso", "grave"):
            filas.append({
                "severidad": c.resultado,
                "tipo": "Campo",
                "campo": NOMBRE_CAMPO.get(c.campo, c.campo),
                "valores": {k: v for k, v in c.valores_por_fuente.items()},
                "valor_final": c.valor_final,
                "sospechosos": [s for s in (c.documento_sospechoso or "").split(",") if s],
                "motivo": c.motivo or "",
                "obj": c,
            })
    for l in r.lineas:
        if l.resultado in ("aviso", "grave"):
            tarifa = l.tarifa
            motivo = "; ".join(l.motivos)
            filas.append({
                "severidad": l.resultado,
                "tipo": "Tarifa" if tarifa and tarifa.estado == "valor_difiere" and len(set(l.valor_por_fuente.values())) <= 1 else "Línea",
                "campo": f"Servicio {l.codigo}" + (f" – {l.descripcion[:60]}" if l.descripcion else ""),
                "valores": {k: f"{v:,.0f}".replace(",", ".") for k, v in l.valor_por_fuente.items()},
                "valor_final": (f"tarifa {tarifa.valor_unitario_tarifa:,.0f}".replace(",", ".") if tarifa and tarifa.valor_unitario_tarifa else None),
                "sospechosos": [s for s in (l.documento_sospechoso or "").split(",") if s],
                "motivo": motivo,
                "obj": l,
            })
    for h in r.hallazgos:
        if h.resultado in ("aviso", "grave"):
            filas.append({
                "severidad": h.resultado,
                "tipo": NOMBRE_HALLAZGO.get(h.tipo, h.tipo),
                "campo": NOMBRE_HALLAZGO.get(h.tipo, h.tipo),
                "valores": {},
                "valor_final": None,
                "sospechosos": [h.documento_sospechoso] if h.documento_sospechoso else [],
                "motivo": h.motivo,
                "obj": h,
            })
    for f in filas:
        f["explicacion"], f["verificar"], f["accion"] = explicar(r, f.pop("obj"))
    orden = {"grave": 0, "aviso": 1}
    return sorted(filas, key=lambda f: orden[f["severidad"]])


def texto_revision(r: ResultadoFactura, filas: list[dict]) -> str:
    lineas = [
        f"FACTURA {r.factura} — {r.estado.upper()}",
        f"Carpeta original: {r.ruta_origen}",
        f"Corrida: {r.corrida_id}",
        "",
        "CAMPOS CON DISCREPANCIA",
        "",
    ]
    presentes = r.documentos_presentes
    archivos = {
        "rips": presentes.get("RIPS"), "xml": presentes.get("XML"), "xml_nota": presentes.get("XML"),
        "pdf_factura": presentes.get("FEV"), "hc": presentes.get("HEV"), "autorizacion_digital": presentes.get("PDE"),
        "autorizacion_scan": presentes.get("PDE"), "cuv": presentes.get("CUV"),
    }
    for f in filas:
        lineas.append(f"[{f['severidad'].upper()}] {f['campo']}")
        lineas.append(f"    QUÉ PASA:        {f['explicacion']}")
        if f["verificar"]:
            lineas.append(f"    CÓMO VERIFICAR:  {f['verificar']}")
        if f["accion"]:
            lineas.append(f"    ACCIÓN SUGERIDA: {f['accion']}")
        if f["valores"]:
            lineas.append("    Valor en cada documento:")
        for fuente, valor in f["valores"].items():
            marca = "   <-- SOSPECHOSO" if fuente in f["sospechosos"] else ""
            archivo = ", ".join(archivos.get(fuente) or [])
            lineas.append(f"    {NOMBRE_FUENTE.get(fuente, fuente):36} {valor!s:34} {archivo}{marca}")
        if f["valor_final"]:
            lineas.append(f"    {'VALOR DE REFERENCIA':36} {f['valor_final']}")
        lineas.append("")
    return "\n".join(lineas).rstrip() + "\n"


# --- copia de carpetas ---------------------------------------------------------


def nombre_eps(raiz: Path) -> str:
    raiz = raiz.resolve()
    return raiz.parent.name if any(p.suffix.lower() in (".xml", ".json", ".pdf") for p in raiz.iterdir() if p.is_file()) else raiz.name


def preparar_destino(destino_raiz: Path, eps: str) -> Path:
    """Crea CON ERROR/<EPS>/. Si ya existe y la generó el validador, se reemplaza completa."""
    carpeta = destino_raiz / eps
    if carpeta.exists():
        if not (carpeta / MARCA).exists():
            raise FileExistsError(f"{carpeta} existe y no fue generada por el validador; no se sobrescribe")
        shutil.rmtree(carpeta)
    carpeta.mkdir(parents=True)
    (carpeta / MARCA).write_text("Carpeta generada automáticamente; se reemplaza en cada corrida.\n", encoding="utf-8")
    return carpeta


def copiar_errores(resultados: list[ResultadoFactura], carpeta_eps: Path) -> int:
    n = 0
    for r in resultados:
        filas = discrepancias(r)
        if not filas or not r.ruta_origen:
            continue
        destino = carpeta_eps / r.factura
        shutil.copytree(r.ruta_origen, destino)  # copia literal: todos los archivos originales
        (destino / "_REVISION.txt").write_text(texto_revision(r, filas), encoding="utf-8")
        (destino / "_resultado.json").write_text(r.model_dump_json(indent=2), encoding="utf-8")
        n += 1
    return n


# --- Excel -----------------------------------------------------------------------

AZUL = "1F3864"
GRIS = "F2F2F2"
ROJO_SUAVE = "F8CBAD"
AMARILLO = "FFE699"
VERDE = "C6EFCE"
ROJO_FUERTE = "FF7C80"
BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def _encabezado(ws, fila: int, titulos: list[str]) -> None:
    for i, t in enumerate(titulos, 1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE
    ws.row_dimensions[fila].height = 32


def _anchos(ws, anchos: dict[int, float]) -> None:
    for col, w in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _color_severidad(celda, severidad: str) -> None:
    celda.fill = PatternFill("solid", fgColor=ROJO_SUAVE if severidad == "grave" else AMARILLO)
    celda.font = Font(bold=True)
    celda.alignment = Alignment(horizontal="center", vertical="top")


def escribir_excel(resultados: list[ResultadoFactura], ruta: Path, eps: str, corrida_id: str, copiadas: set[str]) -> None:
    wb = Workbook()
    todas = {r.factura: discrepancias(r) for r in resultados}
    con_error = [r for r in resultados if todas[r.factura]]

    # Hoja 1: resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = f"Validación de facturas — {eps}"
    ws["A1"].font = Font(bold=True, size=16, color=AZUL)
    datos = [
        ("Corrida", corrida_id),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Facturas procesadas", len(resultados)),
        ("Con discrepancias", len(con_error)),
        ("   graves", sum(1 for r in con_error if any(f["severidad"] == "grave" for f in todas[r.factura]))),
        ("   solo avisos", sum(1 for r in con_error if all(f["severidad"] == "aviso" for f in todas[r.factura]))),
        ("Sin discrepancias", len(resultados) - len(con_error)),
    ]
    for i, (k, v) in enumerate(datos, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=not k.startswith(" "))
        ws.cell(row=i, column=2, value=v)
    fila = len(datos) + 5
    ws.cell(row=fila, column=1, value="Discrepancias por campo").font = Font(bold=True, size=12, color=AZUL)
    _encabezado(ws, fila + 1, ["Campo", "Graves", "Avisos"])
    conteo: dict[str, list[int]] = {}
    for filas in todas.values():
        for f in filas:
            conteo.setdefault(f["campo"] if f["tipo"] == "Campo" else f["tipo"], [0, 0])[0 if f["severidad"] == "grave" else 1] += 1
    for i, (campo, (g, a)) in enumerate(sorted(conteo.items(), key=lambda x: (-x[1][0], -x[1][1])), fila + 2):
        for j, v in enumerate((campo, g, a), 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDE
        if g:
            ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=ROJO_SUAVE)
    _anchos(ws, {1: 34, 2: 26, 3: 12})

    # Hoja 2: discrepancias con el valor de cada fuente
    ws = wb.create_sheet("Discrepancias")
    usadas = [f for f in FUENTES if any(f[0] in d["valores"] for filas in todas.values() for d in filas)]
    titulos = ["Factura", "Severidad", "Tipo", "Campo", "Qué pasa", "Cómo verificarlo", "Acción sugerida"] + [t for _, t, _ in usadas] + ["Valor de referencia", "Documento sospechoso"]
    _encabezado(ws, 1, titulos)
    fila = 2
    filas_ordenadas = sorted(
        ((r, d) for r in con_error for d in todas[r.factura]),
        key=lambda x: (x[1]["severidad"] != "grave", x[0].factura),
    )
    for r, d in filas_ordenadas:
        c = ws.cell(row=fila, column=1, value=r.factura)
        if r.factura in copiadas:
            c.hyperlink = f"{r.factura}/_REVISION.txt"
            c.font = Font(color="0563C1", underline="single")
        _color_severidad(ws.cell(row=fila, column=2, value=d["severidad"].upper()), d["severidad"])
        ws.cell(row=fila, column=3, value=d["tipo"])
        ws.cell(row=fila, column=4, value=d["campo"]).font = Font(bold=True)
        ws.cell(row=fila, column=5, value=d["explicacion"])
        ws.cell(row=fila, column=6, value=d["verificar"] or None)
        ws.cell(row=fila, column=7, value=d["accion"] or None)
        for j, (clave, _, _) in enumerate(usadas, 8):
            valor = d["valores"].get(clave)
            c = ws.cell(row=fila, column=j, value=valor)
            if clave in d["sospechosos"]:
                c.fill = PatternFill("solid", fgColor=ROJO_FUERTE)
                c.font = Font(bold=True)
            elif valor is not None and d["valor_final"] is not None and str(valor) == str(d["valor_final"]):
                c.fill = PatternFill("solid", fgColor=VERDE)
        base = 8 + len(usadas)
        ws.cell(row=fila, column=base, value=d["valor_final"])
        ws.cell(row=fila, column=base + 1, value=", ".join(NOMBRE_FUENTE.get(s, s) for s in d["sospechosos"]) or None)
        for j in range(1, len(titulos) + 1):
            celda = ws.cell(row=fila, column=j)
            celda.border = BORDE
            if j != 2:
                celda.alignment = Alignment(vertical="top", wrap_text=True)
        fila += 1
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(titulos))}{max(fila - 1, 1)}"
    _anchos(ws, {1: 13, 2: 11, 3: 14, 4: 28, 5: 60, 6: 38, 7: 38, **{j: 18 for j in range(8, 8 + len(usadas))},
                 8 + len(usadas): 18, 9 + len(usadas): 26})

    # Hoja 3: todas las facturas
    ws = wb.create_sheet("Facturas")
    _encabezado(ws, 1, ["Factura", "Estado", "Graves", "Avisos", "Fuente principal", "Documentos presentes", "Carpeta original"])
    for i, r in enumerate(sorted(resultados, key=lambda r: (not todas[r.factura], r.factura)), 2):
        g = sum(1 for d in todas[r.factura] if d["severidad"] == "grave")
        a = sum(1 for d in todas[r.factura] if d["severidad"] == "aviso")
        estado = "Con error" if g else ("Con aviso" if a else "OK")
        valores = [r.factura, estado, g, a, r.fuente_principal.upper(), ", ".join(r.documentos_presentes), r.ruta_origen]
        for j, v in enumerate(valores, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = BORDE
        color = ROJO_SUAVE if g else (AMARILLO if a else VERDE)
        ws.cell(row=i, column=2).fill = PatternFill("solid", fgColor=color)
        if r.factura in copiadas:
            ws.cell(row=i, column=1).hyperlink = f"{r.factura}/_REVISION.txt"
            ws.cell(row=i, column=1).font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:G{len(resultados) + 1}"
    _anchos(ws, {1: 13, 2: 12, 3: 8, 4: 8, 5: 15, 6: 34, 7: 70})

    _hoja_ayuda(wb)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)


def _hoja_ayuda(wb: Workbook) -> None:
    ws = wb.create_sheet("Cómo leer este reporte", 1)
    ws["A1"] = "Cómo leer este reporte"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    filas = [
        ("Severidad", ""),
        ("GRAVE", "El dato no coincide entre documentos y afecta el cobro (autorización, documento del paciente, CUPS, valor, CUFE). Corregir antes de radicar."),
        ("AVISO", "Diferencia que hay que revisar pero puede estar bien (nombre con una letra distinta, tipo de consulta, fecha, tarifa de otro contrato)."),
        ("", ""),
        ("Colores en la hoja Discrepancias", ""),
        ("Rojo", "Documento sospechoso: el que tiene el dato distinto."),
        ("Verde", "Documentos que coinciden con el valor de referencia."),
        ("", ""),
        ("Documentos", ""),
        ("RIPS (JSON)", "Archivo FE….json que genera el sistema del hospital. Es la fuente principal."),
        ("XML factura / Nota del XML", "Factura electrónica validada por la DIAN (FE….xml). No se edita: se corrige con nota crédito/débito."),
        ("PDF factura", "Representación gráfica de la factura (FEV_… o FVS_…)."),
        ("Historia clínica", "HEV_…pdf."),
        ("Encabezado de la autorización (hospital)", "Recuadro superior del PDE_…pdf que agrega el hospital al escanear."),
        ("Autorización de la EPS (escaneada)", "La imagen de la autorización emitida por la EPS, leída por OCR. Para el número de autorización es la referencia cuando la lectura es confiable."),
        ("", ""),
        ("Columnas", ""),
        ("Qué pasa", "Explicación del error en lenguaje claro."),
        ("Cómo verificarlo", "Qué archivo abrir y qué buscar para comprobarlo."),
        ("Acción sugerida", "Propuesta de corrección. Confirmar con contabilidad/jurídica antes de aplicarla."),
        ("Valor de referencia", "El valor que el sistema considera correcto (el de la EPS o el de la mayoría de documentos)."),
        ("", ""),
        ("Enlaces", "El número de factura abre el _REVISION.txt de la carpeta copiada, con el mismo detalle."),
    ]
    for i, (a, b) in enumerate(filas, 3):
        ca = ws.cell(row=i, column=1, value=a)
        ws.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
        if a and not b:
            ca.font = Font(bold=True, color=AZUL, size=12)
        elif a:
            ca.font = Font(bold=True)
            if a == "GRAVE":
                ca.fill = PatternFill("solid", fgColor=ROJO_SUAVE)
            elif a == "AVISO":
                ca.fill = PatternFill("solid", fgColor=AMARILLO)
            elif a == "Rojo":
                ca.fill = PatternFill("solid", fgColor=ROJO_FUERTE)
            elif a == "Verde":
                ca.fill = PatternFill("solid", fgColor=VERDE)
    _anchos(ws, {1: 38, 2: 100})
